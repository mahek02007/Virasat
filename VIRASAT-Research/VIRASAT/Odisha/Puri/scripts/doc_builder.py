"""
doc_builder.py - Comprehensive styling and building library for Puri cultural heritage Word documents.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import qn, nsdecls

# Color Palette
COLOR_PRIMARY = RGBColor(139, 30, 15)      # Deep Crimson / Terracotta (#8B1E0F)
COLOR_SECONDARY = RGBColor(184, 115, 51)   # Copper / Ochre (#B87333)
COLOR_GOLD = RGBColor(197, 143, 39)        # Temple Gold (#C58F27)
COLOR_DARK = RGBColor(44, 62, 80)          # Slate Navy (#2C3E50)
COLOR_MUTED = RGBColor(100, 110, 120)      # Muted Gray
COLOR_BODY = RGBColor(33, 37, 41)          # Dark Charcoal

HEX_PRIMARY = "8B1E0F"
HEX_GOLD = "C58F27"
HEX_BG_LIGHT = "FDFBF7"
HEX_BG_CALLOUT = "F7F3E9"
HEX_BORDER = "D1C7B7"
HEX_HEADER_BG = "8B1E0F"
HEX_ZEBRA = "F9F6F0"


def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=120, bottom=120, left=160, right=160):
    """Sets cell internal padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def set_table_borders(table, color="D1C7B7", sz="4", val="single"):
    """Sets elegant subtle borders for a table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)


class PuriDocBuilder:
    """Builder class for creating standardized, richly styled cultural research documents."""

    def __init__(self, title, doc_number="", category="Cultural Heritage Knowledge Base"):
        self.doc = Document()
        self.title = title
        self.doc_number = doc_number
        self.category = category
        self._setup_page_geometry()
        self._setup_styles()
        self._add_header_and_title()

    def _setup_page_geometry(self):
        for section in self.doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)
            section.header_distance = Inches(0.5)
            section.footer_distance = Inches(0.5)

    def _setup_styles(self):
        styles = self.doc.styles
        normal_style = styles['Normal']
        normal_style.font.name = 'Calibri'
        normal_style.font.size = Pt(11)
        normal_style.font.color.rgb = COLOR_BODY
        normal_style.paragraph_format.line_spacing = 1.15
        normal_style.paragraph_format.space_after = Pt(6)

    def _add_header_and_title(self):
        header = self.doc.sections[0].header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"PURI CULTURAL KNOWLEDGE BASE | {self.category.upper()}")
        hrun.font.name = 'Calibri'
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = COLOR_MUTED

        banner_table = self.doc.add_table(rows=1, cols=1)
        banner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        banner_table.autofit = False
        banner_table.columns[0].width = Inches(6.5)
        cell = banner_table.cell(0, 0)
        set_cell_background(cell, HEX_HEADER_BG)
        set_cell_margins(cell, top=200, bottom=200, left=240, right=240)

        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(4)

        if self.doc_number:
            r_num = p.add_run(f"DOCUMENT {self.doc_number} • {self.category.upper()}\n")
            r_num.font.name = 'Calibri'
            r_num.font.size = Pt(9.5)
            r_num.font.bold = True
            r_num.font.color.rgb = COLOR_GOLD

        r_title = p.add_run(self.title)
        r_title.font.name = 'Georgia'
        r_title.font.size = Pt(18)
        r_title.font.bold = True
        r_title.font.color.rgb = RGBColor(255, 255, 255)

        p_sub = cell.add_paragraph()
        p_sub.paragraph_format.space_after = Pt(0)
        r_sub = p_sub.add_run("Puri Cultural Heritage Research Knowledge Base • Odisha, India")
        r_sub.font.name = 'Calibri'
        r_sub.font.size = Pt(10)
        r_sub.font.italic = True
        r_sub.font.color.rgb = RGBColor(240, 230, 220)

        self.doc.add_paragraph().paragraph_format.space_after = Pt(10)

    def add_h1(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Georgia'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return p

    def add_h2(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Georgia'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = COLOR_SECONDARY
        return p

    def add_h3(self, text):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_DARK
        return p

    def add_paragraph(self, text="", bold_prefix=None, italic=False, space_after=6):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15

        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK

        if text:
            r_text = p.add_run(text)
            r_text.font.name = 'Calibri'
            r_text.font.size = Pt(11)
            r_text.font.italic = italic
            r_text.font.color.rgb = COLOR_BODY

        return p

    def add_bullet(self, text, bold_prefix=None):
        p = self.doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15

        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.font.name = 'Calibri'
            r_pre.font.size = Pt(11)
            r_pre.font.bold = True
            r_pre.font.color.rgb = COLOR_DARK

        r_text = p.add_run(text)
        r_text.font.name = 'Calibri'
        r_text.font.size = Pt(11)
        r_text.font.color.rgb = COLOR_BODY
        return p

    def add_callout(self, title, text, tag="FACT", confidence="HIGH"):
        """Adds a standardized callout block with classification tags."""
        tbl = self.doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.5)
        cell = tbl.cell(0, 0)

        set_cell_background(cell, HEX_BG_CALLOUT)
        set_cell_margins(cell, top=140, bottom=140, left=180, right=180)

        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>'
            f'<w:top w:val="none"/>'
            f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{HEX_PRIMARY}"/>'
            f'<w:bottom w:val="none"/>'
            f'<w:right w:val="none"/>'
            f'</w:tcBorders>'
        )
        tcPr.append(borders)

        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(3)
        
        r_tag = p.add_run(f"[{tag.upper()}]  ")
        r_tag.font.name = 'Calibri'
        r_tag.font.size = Pt(9.5)
        r_tag.font.bold = True
        r_tag.font.color.rgb = COLOR_PRIMARY

        if confidence:
            r_conf = p.add_run(f"[CONFIDENCE: {confidence.upper()}]  ")
            r_conf.font.name = 'Calibri'
            r_conf.font.size = Pt(9)
            r_conf.font.bold = True
            r_conf.font.color.rgb = COLOR_SECONDARY

        r_title = p.add_run(title + "\n")
        r_title.font.name = 'Georgia'
        r_title.font.size = Pt(11)
        r_title.font.bold = True
        r_title.font.color.rgb = COLOR_DARK

        p_body = cell.add_paragraph()
        p_body.paragraph_format.space_after = Pt(0)
        p_body.paragraph_format.line_spacing = 1.15
        r_body = p_body.add_run(text)
        r_body.font.name = 'Calibri'
        r_body.font.size = Pt(10.5)
        r_body.font.color.rgb = COLOR_BODY

        self.doc.add_paragraph().paragraph_format.space_after = Pt(4)

    def add_table(self, headers, rows_data, col_widths=None):
        """Adds a clean, professional striped table."""
        tbl = self.doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = False
        set_table_borders(tbl, color="D1C7B7", sz="4")

        hdr_cells = tbl.rows[0].cells
        for idx, header_text in enumerate(headers):
            hdr_cells[idx].text = header_text
            set_cell_background(hdr_cells[idx], HEX_HEADER_BG)
            set_cell_margins(hdr_cells[idx], top=120, bottom=120, left=140, right=140)
            p = hdr_cells[idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Calibri'
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

        for row_idx, row_data in enumerate(rows_data):
            row_cells = tbl.rows[row_idx + 1].cells
            bg_color = HEX_ZEBRA if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, cell_value in enumerate(row_data):
                row_cells[col_idx].text = str(cell_value)
                set_cell_background(row_cells[col_idx], bg_color)
                set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=140, right=140)
                p = row_cells[col_idx].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = 'Calibri'
                    r.font.size = Pt(10)
                    r.font.color.rgb = COLOR_BODY

        if col_widths:
            for i, col in enumerate(tbl.columns):
                w = Inches(col_widths[i]) if i < len(col_widths) else Inches(1.5)
                for cell in col.cells:
                    cell.width = w

        self.doc.add_paragraph().paragraph_format.space_after = Pt(6)
        return tbl

    def add_sources_section(self, sources_list):
        """Standardized sources reference table at document end."""
        self.add_h1("Document References & Sources")
        headers = ["Source ID", "Author / Institution", "Title & Publication", "Type & Tier", "URL / Accession"]
        rows = []
        for s in sources_list:
            rows.append([
                s.get("id", "PUR-S000"),
                s.get("org", "Government of Odisha / ASI"),
                s.get("title", "Official Documentation"),
                f"{s.get('type', 'Report')} (Tier {s.get('tier', '1')})",
                s.get("url", "https://puri.nic.in")
            ])
        self.add_table(headers, rows, [0.9, 1.4, 2.1, 1.1, 1.0])

    def save(self, filepath):
        """Saves the Word document."""
        self.doc.save(filepath)
        print(f"Saved: {filepath}")
