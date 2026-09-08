import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

# Color Palette Constants
COLOR_PRIMARY_HEX = "8B3A1C"      # Bastar Terracotta / Deep Ochre
COLOR_SECONDARY_HEX = "2D5A27"    # Forest Sal Green
COLOR_DARK_HEX = "1A252F"         # Deep Slate Charcoal
COLOR_ACCENT_HEX = "C59B27"       # Antique Bell Metal Gold
COLOR_LIGHT_BG_HEX = "F8F9FA"     # Soft Gray White
COLOR_CALLOUT_BG_HEX = "FDF6E2"   # Warm Parchment
COLOR_MUTED_HEX = "595959"        # Muted Gray

PRIMARY_COLOR = RGBColor(139, 58, 28)
SECONDARY_COLOR = RGBColor(45, 90, 39)
DARK_COLOR = RGBColor(26, 37, 47)
MUTED_COLOR = RGBColor(89, 89, 89)
ACCENT_COLOR = RGBColor(197, 155, 39)

def set_cell_background(cell, hex_color):
    """Sets background color for a table cell."""
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    """Sets cell padding in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color_hex="D3D3D3", sz="4", val="single"):
    """Sets standard clean table borders."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color_hex}"/><w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def create_base_document(title, subtitle=""):
    """Initializes a beautifully styled document with headers, footers, and margins."""
    doc = docx.Document()
    
    # Page setup - 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
        # Header
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("Bastar Heritage Research Initiative | Multidisciplinary Digital Archive")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = MUTED_COLOR
        
        # Footer
        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        frun_left = fp.add_run("Confidential Academic & Cultural Knowledge Base (2026)")
        frun_left.font.name = "Calibri"
        frun_left.font.size = Pt(8.5)
        frun_left.font.color.rgb = MUTED_COLOR
    
    # Configure default style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = DARK_COLOR
    
    # Document Header Title Block
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run(title)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR
    
    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_sub.paragraph_format.space_before = Pt(0)
        p_sub.paragraph_format.space_after = Pt(14)
        run_sub = p_sub.add_run(subtitle)
        run_sub.font.name = "Calibri"
        run_sub.font.size = Pt(13)
        run_sub.font.italic = True
        run_sub.font.color.rgb = SECONDARY_COLOR
        
    return doc

def add_metadata_box(doc, meta_dict):
    """Adds a standardized academic metadata and verification panel."""
    table = doc.add_table(rows=len(meta_dict), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color_hex="C59B27", sz="8")
    
    col_widths = [Inches(2.2), Inches(4.3)]
    for i, (k, v) in enumerate(meta_dict.items()):
        row = table.rows[i]
        
        cell_lbl = row.cells[0]
        cell_lbl.width = col_widths[0]
        set_cell_background(cell_lbl, "FDF8EE")
        set_cell_margins(cell_lbl, top=80, bottom=80, left=120, right=120)
        p0 = cell_lbl.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        p0.paragraph_format.space_before = Pt(0)
        r0 = p0.add_run(k)
        r0.font.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = PRIMARY_COLOR
        
        cell_val = row.cells[1]
        cell_val.width = col_widths[1]
        set_cell_background(cell_val, "FFFFFF")
        set_cell_margins(cell_val, top=80, bottom=80, left=120, right=120)
        p1 = cell_val.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.space_before = Pt(0)
        r1 = p1.add_run(str(v))
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = DARK_COLOR
    
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_after = Pt(12)

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(16)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = PRIMARY_COLOR
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(12)
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = SECONDARY_COLOR
    return h

def add_heading_3(doc, text):
    h = doc.add_paragraph()
    h.paragraph_format.space_before = Pt(8)
    h.paragraph_format.space_after = Pt(2)
    h.paragraph_format.keep_with_next = True
    r = h.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11.5)
    r.font.bold = True
    r.font.color.rgb = DARK_COLOR
    return h

def add_callout(doc, text, title="EVIDENTIARY & ETHNOGRAPHIC NOTE", border_color="8B3A1C", bg_color="FDF6E2"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_color)
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Left border only styling
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    r_t = p.add_run(f"[{title}]\n")
    r_t.font.bold = True
    r_t.font.size = Pt(9.5)
    r_t.font.color.rgb = PRIMARY_COLOR
    
    r_b = p.add_run(text)
    r_b.font.size = Pt(9.5)
    r_b.font.italic = True
    r_b.font.color.rgb = DARK_COLOR
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def add_styled_table(doc, headers, rows_data, col_widths=None):
    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, color_hex="D3D3D3", sz="4")
    
    # Header row
    hdr_cells = table.rows[0].cells
    for j, h_text in enumerate(headers):
        cell = hdr_cells[j]
        if col_widths and j < len(col_widths):
            cell.width = col_widths[j]
        set_cell_background(cell, COLOR_PRIMARY_HEX)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    # Data rows
    for i, row in enumerate(rows_data):
        row_cells = table.rows[i+1].cells
        bg = "FFFFFF" if i % 2 == 0 else "F9FBF9"
        for j, val in enumerate(row):
            cell = row_cells[j]
            if col_widths and j < len(col_widths):
                cell.width = col_widths[j]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.space_before = Pt(0)
            r = p.add_run(str(val))
            r.font.size = Pt(9)
            r.font.color.rgb = DARK_COLOR
            
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

def save_document(doc, filename, output_dir=r"C:\Users\YADAVI\.gemini\antigravity\scratch\BASTAR_CHHATTISGARH_RESEARCH"):
    os.makedirs(output_dir, exist_ok=True)
    filepath = os.path.join(output_dir, filename)
    doc.save(filepath)
    print(f"Generated successfully: {filename}")
    return filepath
